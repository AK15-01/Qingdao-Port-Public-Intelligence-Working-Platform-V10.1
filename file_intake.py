from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re

import pandas as pd

from web_extractor import extract_html


SUPPORTED_EXTENSIONS = {".txt", ".html", ".htm", ".csv", ".docx"}


class UnsupportedFileType(ValueError):
    pass


@dataclass(frozen=True)
class UploadedPublicDocument:
    title: str
    published_date: str
    source_name: str
    text: str
    note: str


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _compact(value: object) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def parse_uploaded_public_file(filename: str, content: bytes) -> UploadedPublicDocument:
    suffix = Path(filename or "").suffix.casefold()
    if suffix not in SUPPORTED_EXTENSIONS:
        if suffix == ".pdf":
            raise UnsupportedFileType("当前版本不解析PDF，请在阅读原文后粘贴事实摘要。")
        if suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}:
            raise UnsupportedFileType("当前版本不识别图片文字，请在阅读原图后手工粘贴。")
        raise UnsupportedFileType("只支持 TXT、HTML、CSV 或 DOCX 公共资料文件。")
    if not content:
        raise ValueError("上传文件为空。")

    if suffix == ".txt":
        text = _decode_text(content).strip()
        title = next((line.strip() for line in text.splitlines() if line.strip()), "")[:200]
        return UploadedPublicDocument(title, "", "", text[:30000], "已读取TXT文本，标题和来源需人工核对。")

    if suffix in {".html", ".htm"}:
        html = _decode_text(content)
        result = extract_html(
            html,
            requested_url="https://local-upload.invalid/document",
            final_url="https://local-upload.invalid/document",
        )
        return UploadedPublicDocument(
            result.title,
            result.published_date,
            result.source_name,
            result.text,
            "已从本地HTML提取正文；本地文件地址不会作为原文链接入库，需补充公开URL。",
        )

    if suffix == ".csv":
        dataframe = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        if dataframe.empty:
            raise ValueError("CSV没有可读取的数据行。")
        row = dataframe.iloc[0]
        title = _compact(row.get("title") or row.get("fetched_title"))
        published_date = _compact(row.get("event_date") or row.get("fetched_date"))
        source_name = _compact(row.get("source_name") or row.get("fetched_source_name"))
        text = _compact(
            row.get("fetched_text")
            or row.get("summary")
            or row.get("text")
            or row.get("content")
        )
        if not text:
            text = "；".join(
                f"{column}：{_compact(value)}" for column, value in row.items() if _compact(value)
            )
        note = "CSV只读取第一条记录，不执行批量导入。"
        if len(dataframe) > 1:
            note += f" 文件共有 {len(dataframe)} 行，其余行未处理。"
        return UploadedPublicDocument(title, published_date, source_name, text[:30000], note)

    from docx import Document

    document = Document(BytesIO(content))
    paragraphs = [_compact(paragraph.text) for paragraph in document.paragraphs]
    text = "\n".join(item for item in paragraphs if item)
    title = next((item for item in paragraphs if item), "")[:200]
    return UploadedPublicDocument(
        title,
        "",
        "",
        text[:30000],
        "已读取DOCX正文；发布日期、来源和原文URL需人工补充并核对。",
    )
