from __future__ import annotations

import json
from pathlib import Path
from typing import Mapping, Optional

from openpyxl import load_workbook

from platform_db import connect, initialize_database, new_id, now_iso, transaction
from review_provenance import HUMAN_REVIEWER_TYPES


ROOT = Path(__file__).resolve().parent
LABEL_FIELDS = {
    "正确标题",
    "正确发布日期",
    "正确发布机构",
    "正文是否合格",
    "是否包含乱码",
    "是否包含导航噪声",
    "正确类别",
    "事实摘要是否准确",
    "evidence_quotes是否存在于原文",
    "潜在影响是否合理",
    "是否具有业务价值",
    "是否允许进入内部研究版",
    "是否允许进入客户交付版",
    "人工标注状态",
    "人工备注",
}
AUDIT_FIELDS = {
    "reviewer_type",
    "reviewer_name",
    "reviewed_at",
    "review_method",
    "review_version",
    "reviewer_note",
}
MAJOR_FACT_FIELDS = {
    "title",
    "event_date",
    "summary",
    "affected_area",
    "category",
    "status",
}


def load_review_progress(qa_db_path: Path, workspace_id: str) -> dict[str, object]:
    """Load UI progress only; progress never counts as a completed review."""

    initialize_database(qa_db_path)
    with connect(qa_db_path) as connection:
        row = connection.execute(
            "SELECT * FROM qa_review_progress WHERE workspace_id=?",
            (workspace_id,),
        ).fetchone()
    if not row:
        return {
            "workspace_id": workspace_id,
            "current_document_id": "",
            "filter_status": "未审核",
            "reviewer_type": "human_user",
            "reviewer_name": "",
            "draft": {},
        }
    result = dict(row)
    try:
        result["draft"] = json.loads(str(result.get("draft_json") or "{}"))
    except json.JSONDecodeError:
        result["draft"] = {}
    return result


def save_review_progress(
    qa_db_path: Path,
    workspace_id: str,
    *,
    current_document_id: str,
    filter_status: str,
    reviewer_type: str = "human_user",
    reviewer_name: str = "",
    draft: Optional[Mapping[str, object]] = None,
) -> None:
    """Persist a resumable cursor/draft without changing review provenance."""

    if reviewer_type not in HUMAN_REVIEWER_TYPES:
        raise ValueError("审核进度只允许记录项目使用者或行业复核人身份")
    timestamp = now_iso()
    with transaction(qa_db_path) as connection:
        connection.execute(
            """INSERT INTO qa_review_progress(
            workspace_id,current_document_id,filter_status,reviewer_type,
            reviewer_name,draft_json,updated_at
            ) VALUES(?,?,?,?,?,?,?)
            ON CONFLICT(workspace_id) DO UPDATE SET
            current_document_id=excluded.current_document_id,
            filter_status=excluded.filter_status,
            reviewer_type=excluded.reviewer_type,
            reviewer_name=excluded.reviewer_name,
            draft_json=excluded.draft_json,updated_at=excluded.updated_at""",
            (
                workspace_id,
                current_document_id,
                filter_status,
                reviewer_type,
                reviewer_name.strip(),
                json.dumps(dict(draft or {}), ensure_ascii=False, default=str),
                timestamp,
            ),
        )


def qa_review_status(record: Mapping[str, object]) -> str:
    reviewer_type = str(record.get("reviewer_type") or "")
    if reviewer_type not in HUMAN_REVIEWER_TYPES:
        return "未审核"
    state = str(record.get("人工标注状态") or "")
    if state in {"已完成", "修改后完成"}:
        return "接受" if state == "已完成" else "修改后接受"
    if state == "已驳回":
        return "驳回"
    if state == "需要第二人复核":
        return "待二审"
    if state == "正文有问题":
        return "正文有问题"
    return "未审核"


def latest_qa_database(root: Path = ROOT) -> Optional[Path]:
    candidates = sorted(
        (root / "output" / "qa_real").glob("run-*/portscope_qa.db"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return candidates[0] if candidates else None


def qa_context(
    *,
    labels_path: Path = ROOT / "qa" / "real_acceptance_labels.xlsx",
    qa_db_path: Optional[Path] = None,
) -> dict[str, object] | None:
    qa_db = Path(qa_db_path) if qa_db_path else latest_qa_database()
    if not qa_db or not qa_db.exists() or not labels_path.exists():
        return None
    initialize_database(qa_db)
    workbook = load_workbook(labels_path, read_only=True, data_only=True)
    try:
        sheet = workbook["人工标注"]
        headers = [str(cell.value or "") for cell in sheet[1]]
        rows = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
            if str(row.get("document_id") or "").strip():
                rows.append(row)
    finally:
        workbook.close()
    with connect(qa_db) as connection:
        workspace = connection.execute(
            "SELECT workspace_id,workspace_name FROM workspaces ORDER BY created_at LIMIT 1"
        ).fetchone()
        documents = {
            str(row["document_id"]): dict(row)
            for row in connection.execute(
                """SELECT d.*,e.event_id,e.title AS event_title,e.summary,e.impact,
                e.event_date,e.affected_area,e.affected_period,e.involved_entities,
                e.category,e.status,e.risk_level,e.human_verified,
                e.evidence_verified,e.reviewer_type AS event_reviewer_type
                FROM documents d LEFT JOIN events e ON e.document_id=d.document_id
                AND e.record_state='active' WHERE d.is_current=1"""
            ).fetchall()
        }
    for row in rows:
        row["document"] = documents.get(str(row["document_id"]), {})
    return {
        "db_path": qa_db,
        "labels_path": labels_path,
        "workspace_id": str(workspace["workspace_id"]) if workspace else "",
        "workspace_name": str(workspace["workspace_name"]) if workspace else "QA工作空间",
        "records": rows,
    }


def _write_updated_workbook(
    labels_path: Path,
    document_id: str,
    updates: Mapping[str, object],
) -> Path:
    workbook = load_workbook(labels_path)
    sheet = workbook["人工标注"]
    headers = {str(cell.value or ""): cell.column for cell in sheet[1]}
    target_row = 0
    document_column = headers.get("document_id")
    if not document_column:
        workbook.close()
        raise ValueError("人工标注表缺少 document_id 列")
    for row_index in range(2, sheet.max_row + 1):
        if str(sheet.cell(row_index, document_column).value or "") == document_id:
            target_row = row_index
            break
    if not target_row:
        workbook.close()
        raise KeyError("人工标注表中不存在该文档")
    for field, value in updates.items():
        if field not in headers:
            headers[field] = sheet.max_column + 1
            sheet.cell(1, headers[field], field)
        sheet.cell(target_row, headers[field], value)
    temporary = labels_path.with_suffix(labels_path.suffix + ".tmp")
    workbook.save(temporary)
    workbook.close()
    return temporary


def save_qa_label_review(
    document_id: str,
    *,
    qa_db_path: Path,
    labels_path: Path,
    reviewer_type: str,
    reviewer_name: str,
    decision: str,
    edits: Mapping[str, object],
    event_edits: Optional[Mapping[str, object]] = None,
    checklist: Mapping[str, object],
    reviewer_note: str = "",
) -> dict[str, object]:
    if reviewer_type not in HUMAN_REVIEWER_TYPES:
        raise ValueError("只有项目使用者或行业复核人可以保存独立真人核验")
    if not reviewer_name.strip():
        raise ValueError("必须填写核验人姓名或内部代号")
    if decision not in {"accept", "accept_modified", "reject", "body_issue", "needs_second_review"}:
        raise ValueError("不支持的审核决定")
    if decision in {"accept", "accept_modified"} and not (
        checklist.get("source_opened") and checklist.get("facts_checked")
    ):
        raise ValueError("接受前必须打开原文并逐项核对事实")
    context = qa_context(labels_path=labels_path, qa_db_path=qa_db_path)
    if not context:
        raise FileNotFoundError("QA数据库或人工标签文件不存在")
    record = next(
        (
            item
            for item in context["records"]
            if str(item.get("document_id") or "") == document_id
        ),
        None,
    )
    if not record:
        raise KeyError("QA标签不存在")
    document = dict(record.get("document") or {})
    event_id = str(document.get("event_id") or "")
    before = {
        key: value
        for key, value in record.items()
        if key != "document"
    }
    timestamp = now_iso()
    event_updates = {
        key: ("" if value is None else str(value).strip())
        for key, value in dict(event_edits or {}).items()
    }
    major_changes = [
        key
        for key in MAJOR_FACT_FIELDS
        if key in event_updates
        and str(document.get(key if key != "title" else "event_title") or "")
        != str(event_updates[key] or "")
    ]
    if decision == "accept" and major_changes:
        raise ValueError("检测到重大事实字段变化，请选择“修改后接受”并填写修改原因")
    if decision == "accept_modified" and major_changes and not reviewer_note.strip():
        raise ValueError("修改标题、日期、摘要、地域、类别或状态等重大事实时必须填写修改原因")
    updated = {
        key: value
        for key, value in edits.items()
        if key in LABEL_FIELDS
    }
    updated.update(
        {
            "人工标注状态": {
                "accept": "已完成",
                "accept_modified": "修改后完成",
                "reject": "已驳回",
                "body_issue": "正文有问题",
                "needs_second_review": "需要第二人复核",
            }[decision],
            "人工备注": reviewer_note,
            "reviewer_type": reviewer_type,
            "reviewer_name": reviewer_name.strip(),
            "reviewed_at": timestamp,
            "review_method": "工作台逐项对照原文",
            "review_version": "0.5.0-beta",
            "reviewer_note": reviewer_note,
        }
    )
    changed = [
        key for key, value in updated.items()
        if str(before.get(key) or "") != str(value or "")
    ]
    temporary = _write_updated_workbook(labels_path, document_id, updated)
    workspace_id = str(context["workspace_id"])
    try:
        if event_id:
            from review_service import save_event_review

            save_event_review(
                event_id,
                workspace_id,
                qa_db_path,
                decision=decision,
                reviewer_type=reviewer_type,
                reviewer_name=reviewer_name,
                reviewer_note=reviewer_note,
                edits=event_updates,
                checklist=checklist,
            )
        with transaction(qa_db_path) as connection:
            connection.execute(
                """INSERT INTO qa_label_reviews(
                review_id,workspace_id,document_id,event_id,reviewer_type,reviewer_name,
                reviewed_at,review_method,review_version,reviewer_note,decision,
                checklist_json,before_json,after_json,changed_fields_json,created_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    new_id("QREV"),
                    workspace_id,
                    document_id,
                    event_id,
                    reviewer_type,
                    reviewer_name.strip(),
                    timestamp,
                    "工作台逐项对照原文",
                    "0.5.0-beta",
                    reviewer_note,
                    decision,
                    json.dumps(dict(checklist), ensure_ascii=False, default=str),
                    json.dumps(before, ensure_ascii=False, default=str),
                    json.dumps({**before, **updated}, ensure_ascii=False, default=str),
                    json.dumps(changed, ensure_ascii=False),
                    timestamp,
                ),
            )
        temporary.replace(labels_path)
        save_review_progress(
            qa_db_path,
            workspace_id,
            current_document_id="",
            filter_status="未审核",
            reviewer_type=reviewer_type,
            reviewer_name=reviewer_name,
            draft={},
        )
    except Exception:
        if temporary.exists():
            temporary.unlink()
        raise
    return {
        "document_id": document_id,
        "event_id": event_id,
        "reviewer_type": reviewer_type,
        "decision": decision,
        "changed_fields": changed,
        "reviewed_at": timestamp,
    }
