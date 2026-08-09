from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
import pytest

from event_pipeline import batch_confirm_events
from platform_report import generate_platform_report
from tests.test_deepseek_and_rag import _store
from workspace_store import workspace_paths


def _sheet_names(path: Path):
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/workbook.xml"))
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    return [item.attrib["name"] for item in root.findall("x:sheets/x:sheet", ns)]


def test_platform_report_uses_only_verified_licensed_data_and_has_seven_sheets(tmp_path: Path):
    db, root, ws, _, _, event_id, _ = _store(tmp_path)
    paths = workspace_paths(ws["workspace_id"], root)
    batch_confirm_events(
        [event_id], ws["workspace_id"], db,
        reviewer_type="human_user", reviewer_name="测试人员",
        review_method="测试中逐项核对",
    )
    artifacts = generate_platform_report(ws, paths, client=None, start_date="2026-07-01", end_date="2026-07-31",
        report_title="SQLite商业报告", analyst="测试分析", db_path=db)
    assert artifacts.event_ids == (event_id,)
    assert _sheet_names(artifacts.xlsx_path) == ["事件明细", "当前风险", "已解除风险", "商机清单", "来源清单", "数据质量", "抓取运行记录"]
    assert Document(artifacts.docx_path).paragraphs
    assert "https://example.com/a" in artifacts.html_path.read_text(encoding="utf-8")


def test_unverified_event_is_not_in_platform_report(tmp_path: Path):
    db, root, ws, _, _, event_id, _ = _store(tmp_path)
    paths = workspace_paths(ws["workspace_id"], root)
    artifacts = generate_platform_report(ws, paths, client=None, start_date="2026-07-01", end_date="2026-07-31",
        report_title="空报告", analyst="测试分析", db_path=db)
    assert event_id not in artifacts.event_ids


def test_unknown_report_mode_cannot_bypass_empty_report_gate(tmp_path: Path):
    db, root, ws, _, _, _, _ = _store(tmp_path)
    paths = workspace_paths(ws["workspace_id"], root)
    with pytest.raises(ValueError, match="不支持的报告模式"):
        generate_platform_report(
            ws,
            paths,
            client=None,
            start_date="2026-07-01",
            end_date="2026-07-31",
            report_title="未知模式报告",
            analyst="测试分析",
            db_path=db,
            report_mode="unknown-mode",
        )
