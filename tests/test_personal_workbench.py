from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from evidence_binding import (
    invalidate_superseded_document_evidence,
    locate_evidence_quote,
    replace_event_evidence,
    suggest_evidence_candidates,
)
from backup_workspace import backup_workspace, restore_database_snapshot, sanitize_backup_archive
from evidence_repair import list_evidence_issues
from operation_store import (
    clear_home_operations,
    close_operation,
    finish_operation,
    get_operation,
    home_operations,
    list_operations,
    start_operation,
    transition_operation,
)
from platform_db import connect, initialize_database, table_csv_bytes, transaction
from platform_report import eligible_events_dataframe
from rag.keyword_search import KeywordSearcher
from review_service import get_review_record, save_event_review
from qa_review_store import save_qa_label_review
from tests.test_deepseek_and_rag import _store
from workspace_store import create_workspace
from local_import_pipeline import import_public_material


def _verified_event(tmp_path: Path):
    db, root, workspace, source_id, document_id, event_id, event = _store(tmp_path)
    with transaction(db) as connection:
        bindings = replace_event_evidence(
            connection,
            event_id,
            document_id,
            ["公开测试信息显示相关水域临时限航"],
        )
    assert bindings and bindings[0].verification_status == "已验证"
    return db, root, workspace, source_id, document_id, event_id, event


def test_completed_task_closes_from_home_but_remains_in_history(tmp_path: Path):
    db = tmp_path / "history.db"
    workspace = create_workspace({"workspace_name": "历史测试"}, db, tmp_path / "data")
    operation_id = start_operation(
        workspace["workspace_id"],
        "crawl",
        db,
        ui_session_id="SESSION-A",
        input_summary="小规模采集",
    )
    finish_operation(
        operation_id,
        db,
        counts={"documents_created": 2, "skipped_count": 1},
        result_summary="完成",
    )
    assert [item["operation_run_id"] for item in home_operations(workspace["workspace_id"], "SESSION-A", db)] == [operation_id]
    close_operation(operation_id, workspace["workspace_id"], db)
    assert home_operations(workspace["workspace_id"], "SESSION-A", db) == []
    history = list_operations(workspace["workspace_id"], db)
    saved = next(item for item in history if item["operation_run_id"] == operation_id)
    assert saved["status"] == "succeeded"
    assert saved["is_archived"] == 1
    assert saved["home_visible"] == 0
    assert saved["archived_at"]
    assert saved["documents_created"] == 2


def test_clearing_home_never_deletes_business_history(tmp_path: Path):
    db = tmp_path / "clear.db"
    workspace = create_workspace({"workspace_name": "清理测试"}, db, tmp_path / "data")
    operation_id = start_operation(
        workspace["workspace_id"],
        "system_check",
        db,
        ui_session_id="SESSION-A",
    )
    finish_operation(operation_id, db)
    assert clear_home_operations(workspace["workspace_id"], "SESSION-A", db) == 1
    assert get_operation(operation_id, db) is not None
    assert operation_id in {
        item["operation_run_id"]
        for item in list_operations(workspace["workspace_id"], db)
    }


def test_operation_state_transitions_are_strict(tmp_path: Path):
    db = tmp_path / "state.db"
    workspace = create_workspace({"workspace_name": "状态测试"}, db, tmp_path / "data")
    operation_id = start_operation(workspace["workspace_id"], "file_import", db)
    with pytest.raises(ValueError):
        transition_operation(operation_id, "queued", db)
    finish_operation(operation_id, db, status="partially_succeeded", error_summary="一项失败")
    with pytest.raises(ValueError):
        transition_operation(operation_id, "running", db)


def test_human_modified_review_keeps_before_after_diff(tmp_path: Path):
    db, _, workspace, _, _, event_id, _ = _verified_event(tmp_path)
    result = save_event_review(
        event_id,
        workspace["workspace_id"],
        db,
        decision="accept_modified",
        reviewer_type="human_user",
        reviewer_name="项目使用者",
        edits={"title": "人工核对后的限航通知", "manual_risk_level": "中"},
        checklist={"source_opened": True, "facts_checked": True},
    )
    assert result["human_verified"] is True
    assert set(result["changed_fields"]) == {"title", "manual_risk_level"}
    record = get_review_record(event_id, workspace["workspace_id"], db)
    audit = record["reviews"][0]
    assert json.loads(audit["before_json"])["title"] == "测试航行警告"
    assert json.loads(audit["after_json"])["title"] == "人工核对后的限航通知"


@pytest.mark.parametrize(
    ("decision", "expected_state", "second_review"),
    [
        ("reject", "rejected", 0),
        ("body_issue", "quarantined", 0),
        ("needs_second_review", "active", 1),
    ],
)
def test_reject_body_issue_and_second_review_never_enter_customer_report(
    tmp_path: Path,
    decision: str,
    expected_state: str,
    second_review: int,
):
    db, _, workspace, _, _, event_id, _ = _verified_event(tmp_path)
    save_event_review(
        event_id,
        workspace["workspace_id"],
        db,
        decision=decision,
        reviewer_type="human_user",
        reviewer_name="项目使用者",
        checklist={"source_opened": True, "facts_checked": True},
    )
    with connect(db) as connection:
        event = connection.execute(
            "SELECT record_state,requires_second_review,report_eligible FROM events WHERE event_id=?",
            (event_id,),
        ).fetchone()
    assert event["record_state"] == expected_state
    assert event["requires_second_review"] == second_review
    assert event["report_eligible"] == 0
    assert eligible_events_dataframe(
        workspace["workspace_id"], db, report_mode="客户交付版"
    ).empty


def test_evidence_gate_rejects_punctuation_rewrite_and_stitched_text():
    source = "第一段事实明确。第二段事实完整。"
    assert locate_evidence_quote("第一段事实明确。", source).verification_status == "已验证"
    assert locate_evidence_quote("第一段事实明确！", source).verification_status == "未定位"
    assert locate_evidence_quote("第一段事实明确第二段事实完整", source).verification_status == "未定位"
    candidates = suggest_evidence_candidates("第一段事实清楚。", source)
    assert candidates
    assert all(item["verification_status"] == "仅供人工候选，不得自动认证" for item in candidates)


def test_document_version_change_invalidates_verified_binding(tmp_path: Path):
    db, _, _, _, document_id, event_id, _ = _verified_event(tmp_path)
    with transaction(db) as connection:
        assert invalidate_superseded_document_evidence(connection, document_id) == 1
    with connect(db) as connection:
        evidence = connection.execute(
            "SELECT verification_status FROM event_evidence WHERE event_id=?",
            (event_id,),
        ).fetchone()
        event = connection.execute(
            "SELECT evidence_verified,report_eligible FROM events WHERE event_id=?",
            (event_id,),
        ).fetchone()
    assert evidence["verification_status"] == "版本失配"
    assert tuple(event) == (0, 0)


def test_quarantined_data_is_excluded_from_rag_and_reports(tmp_path: Path):
    db, _, workspace, _, document_id, event_id, _ = _verified_event(tmp_path)
    with transaction(db) as connection:
        connection.execute(
            "UPDATE documents SET record_state='quarantined',quarantine_reason='测试隔离' WHERE document_id=?",
            (document_id,),
        )
        connection.execute(
            "UPDATE events SET record_state='quarantined',report_eligible=0 WHERE event_id=?",
            (event_id,),
        )
    assert KeywordSearcher(db).search("临时限航", workspace["workspace_id"]) == []
    assert eligible_events_dataframe(
        workspace["workspace_id"], db, report_mode="内部研究版"
    ).empty


def test_internal_analysis_is_not_blocked_by_unknown_reuse_but_fulltext_export_is(
    tmp_path: Path,
):
    db, _, workspace, source_id, _, event_id, _ = _verified_event(tmp_path)
    with transaction(db) as connection:
        connection.execute(
            """UPDATE sources SET commercial_reuse_status='未明确',
            internal_analysis_allowed=1,customer_summary_allowed=1,short_quote_allowed=1,
            fulltext_redistribution_allowed=0,raw_data_resale_allowed=0
            WHERE source_id=?""",
            (source_id,),
        )
    internal = eligible_events_dataframe(
        workspace["workspace_id"], db, human_verified_only=False, report_mode="内部研究版"
    )
    assert event_id in set(internal["event_id"])
    with pytest.raises(PermissionError, match="二次确认"):
        table_csv_bytes(
            "documents",
            workspace["workspace_id"],
            db,
            export_mode="fulltext",
        )
    with pytest.raises(PermissionError, match="未允许全文再分发"):
        table_csv_bytes(
            "documents",
            workspace["workspace_id"],
            db,
            export_mode="fulltext",
            confirmed=True,
        )


@pytest.mark.parametrize("folder", ["中文路径", "space path", "hash#path", "and&path", "paren(path)"])
def test_sqlite_initialization_supports_windows_special_paths(tmp_path: Path, folder: str):
    db = tmp_path / folder / "工作台 数据#.db"
    initialize_database(db)
    initialize_database(db)
    with connect(db) as connection:
        assert connection.execute("PRAGMA quick_check").fetchone()[0] == "ok"
        assert connection.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='operation_runs'"
        ).fetchone()[0] == 1


def test_current_five_historical_evidence_issues_have_a_non_verified_repair_record():
    report = Path(__file__).resolve().parents[1] / "qa" / "evidence_repair_report.json"
    payload = json.loads(report.read_text(encoding="utf-8"))
    assert payload["issue_count"] == 5
    assert all(item["verification_status"] != "已验证" for item in payload["issues"])
    assert all(item["reason"] and item["suggested_action"] for item in payload["issues"])
    assert any(item["meaning_change"] for item in payload["issues"])


def test_qa_human_review_preserves_agent_label_in_audit_chain(tmp_path: Path):
    db, _, workspace, _, document_id, event_id, _ = _verified_event(tmp_path)
    labels = tmp_path / "labels.xlsx"
    workbook = Workbook()
    workbook.active.title = "填写说明"
    sheet = workbook.create_sheet("人工标注")
    headers = [
        "document_id", "原文URL", "正确标题", "正确发布日期", "正确发布机构",
        "正文是否合格", "是否包含乱码", "是否包含导航噪声", "正确类别",
        "事实摘要是否准确", "evidence_quotes是否存在于原文", "潜在影响是否合理",
        "是否具有业务价值", "是否允许进入内部研究版", "是否允许进入客户交付版",
        "人工标注状态", "人工备注", "reviewer_type", "reviewer_name", "reviewed_at",
        "review_method", "review_version", "reviewer_note",
    ]
    sheet.append(headers)
    sheet.append(
        [
            document_id, "https://example.com/a", "测试航行警告", "2026-07-20",
            "测试公开机构", "是", "否", "否", "航行警告", "是", "是", "是",
            "高", "是", "否", "自动标注", "", "codex_agent", "Codex",
            "2026-07-20T10:00:00+08:00", "自动验收", "0.3", "",
        ]
    )
    workbook.save(labels)
    workbook.close()
    result = save_qa_label_review(
        document_id,
        qa_db_path=db,
        labels_path=labels,
        reviewer_type="human_user",
        reviewer_name="项目使用者",
        decision="accept",
        edits={"人工标注状态": "已完成", "正确标题": "测试航行警告"},
        event_edits={},
        checklist={"source_opened": True, "facts_checked": True},
        reviewer_note="已逐项核对。",
    )
    assert result["event_id"] == event_id
    reloaded = load_workbook(labels, read_only=True, data_only=True)
    headers = [cell.value for cell in reloaded["人工标注"][1]]
    values = [cell.value for cell in reloaded["人工标注"][2]]
    row = dict(zip(headers, values))
    reloaded.close()
    assert row["reviewer_type"] == "human_user"
    with connect(db) as connection:
        audit = connection.execute(
            "SELECT before_json,after_json FROM qa_label_reviews WHERE document_id=?",
            (document_id,),
        ).fetchone()
    assert json.loads(audit["before_json"])["reviewer_type"] == "codex_agent"
    assert json.loads(audit["after_json"])["reviewer_type"] == "human_user"


def test_database_backup_restore_to_explicit_special_path(tmp_path: Path):
    source = tmp_path / "源 数据#.db"
    initialize_database(source)
    archive = tmp_path / "backup &(1).zip"
    with ZipFile(archive, "w") as output:
        output.write(source, "data/portscope.db")
    target = tmp_path / "恢复 #空间" / "正式 数据.db"
    with pytest.raises(PermissionError):
        restore_database_snapshot(archive, target)
    result = restore_database_snapshot(archive, target, confirmed=True)
    assert result["quick_check"] == "ok"
    with connect(target) as connection:
        assert connection.execute("PRAGMA quick_check").fetchone()[0] == "ok"


def test_workspace_backup_and_sanitizer_exclude_env_and_pytest_output(
    tmp_path: Path,
    monkeypatch,
):
    import backup_workspace as backup_module

    root = tmp_path / "工作台 #备份"
    (root / "config").mkdir(parents=True)
    (root / "output" / "pytest-run" / "case").mkdir(parents=True)
    (root / "data").mkdir(parents=True)
    (root / "config" / "settings.json").write_text("{}", encoding="utf-8")
    (root / "output" / "pytest-run" / "case" / ".env").write_text(
        "DEEPSEEK_API_KEY=TEST_PLACEHOLDER",
        encoding="utf-8",
    )
    (root / ".env.example").write_text("DEEPSEEK_API_KEY=", encoding="utf-8")
    monkeypatch.setattr(backup_module, "ROOT", root)
    archive = tmp_path / "备份 #(1).zip"
    result = backup_workspace(archive)
    assert result["file_count"] == 2
    with ZipFile(archive) as payload:
        names = payload.namelist()
    assert "config/settings.json" in names
    assert ".env.example" in names
    assert not any(Path(name).name == ".env" for name in names)
    assert not any("pytest-" in name for name in names)

    unsafe = tmp_path / "旧备份.zip"
    with ZipFile(unsafe, "w") as payload:
        payload.writestr("data/portscope.db", b"safe")
        payload.writestr("output/pytest-old/case/.env", b"secret")
    cleaned = sanitize_backup_archive(unsafe)
    assert cleaned["removed"] == 1
    with ZipFile(unsafe) as payload:
        assert payload.namelist() == ["data/portscope.db"]


def test_single_public_text_import_uses_sqlite_and_unified_history(tmp_path: Path):
    db = tmp_path / "导入 数据#.db"
    root = tmp_path / "导入 根目录"
    workspace = create_workspace({"workspace_name": "导入测试"}, db, root)
    result = import_public_material(
        workspace["workspace_id"],
        db,
        root,
        source_url="https://example.com/public/notice",
        title="公开限航测试通知",
        published_at="2026-07-27",
        source_name="测试公开机构",
        pasted_text=(
            "测试公开机构发布临时限航通知。相关船舶应核对原文中的时间和水域范围。"
            "本信息仅用于自动化测试，不代表任何真实港口运行情况。"
        ),
        ui_session_id="SESSION-IMPORT",
    )
    assert result["document_id"] and result["event_id"]
    with connect(db) as connection:
        document = connection.execute(
            "SELECT record_state,title FROM documents WHERE document_id=?",
            (result["document_id"],),
        ).fetchone()
        operation = connection.execute(
            "SELECT status,documents_created,events_created FROM operation_runs WHERE operation_run_id=?",
            (result["operation_run_id"],),
        ).fetchone()
    assert document["record_state"] == "active"
    assert tuple(operation) == ("succeeded", 1, 1)
