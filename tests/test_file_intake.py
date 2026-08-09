from io import BytesIO

import pytest
from docx import Document

from file_intake import UnsupportedFileType, parse_uploaded_public_file


def test_txt_upload_is_parsed():
    result = parse_uploaded_public_file("notice.txt", "测试公告\n这是公开正文内容。".encode("utf-8"))
    assert result.title == "测试公告"
    assert "公开正文" in result.text


def test_html_upload_uses_existing_extractor():
    html = b"<html><head><title>HTML notice</title></head><body><article><h1>HTML notice</h1><p>Long enough public body content for deterministic extraction and review.</p></article></body></html>"
    result = parse_uploaded_public_file("notice.html", html)
    assert result.title == "HTML notice"
    assert "public body content" in result.text


def test_csv_upload_reads_only_first_row():
    content = "title,summary,source_name\n第一条,第一条公开摘要,公开来源\n第二条,不应处理,公开来源\n".encode("utf-8")
    result = parse_uploaded_public_file("notice.csv", content)
    assert result.title == "第一条"
    assert "只读取第一条" in result.note
    assert "第二条" not in result.text


def test_docx_upload_can_be_read():
    document = Document()
    document.add_paragraph("DOCX公开公告")
    document.add_paragraph("这是需要人工核对的公开正文。")
    buffer = BytesIO()
    document.save(buffer)
    result = parse_uploaded_public_file("notice.docx", buffer.getvalue())
    assert result.title == "DOCX公开公告"
    assert "公开正文" in result.text


def test_pdf_and_images_are_rejected_with_clear_message():
    with pytest.raises(UnsupportedFileType, match="不解析PDF"):
        parse_uploaded_public_file("notice.pdf", b"%PDF")
    with pytest.raises(UnsupportedFileType, match="不识别图片"):
        parse_uploaded_public_file("notice.png", b"png")
