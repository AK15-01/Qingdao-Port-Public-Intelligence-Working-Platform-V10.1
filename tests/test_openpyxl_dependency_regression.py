from __future__ import annotations

import importlib
from pathlib import Path
import subprocess

from openpyxl import Workbook, load_workbook

import scripts.launch_portscope as launcher
from scripts.audit_runtime_dependencies import audit_runtime_dependencies
from qa_review_store import _write_updated_workbook, qa_context
import ui_daily
from platform_db import initialize_database
from workspace_store import create_workspace


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROJECT_PYTHON = PROJECT_ROOT / ".venv" / "Scripts" / "python.exe"


def test_openpyxl_is_a_declared_launcher_core_dependency():
    assert "openpyxl" in launcher.CORE_MODULES
    assert launcher.CORE_PACKAGE_NAMES["openpyxl"] == "openpyxl"
    assert importlib.import_module("openpyxl").__version__ == "3.1.5"


def test_check_only_succeeds_with_openpyxl_in_project_environment():
    if not PROJECT_PYTHON.exists():
        return
    result = subprocess.run(
        [
            str(PROJECT_PYTHON),
            str(PROJECT_ROOT / "scripts" / "launch_portscope.py"),
            "--check-only",
            "--skip-port-check",
        ],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=60,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "核心启动条件通过" in result.stdout


def test_missing_openpyxl_returns_nonzero_and_names_the_module(monkeypatch, capsys):
    real_status = launcher._module_status

    def fake_status(module_name: str, *, import_module: bool = True) -> str:
        if module_name == "openpyxl":
            return "不可用：ModuleNotFoundError: No module named 'openpyxl'"
        return real_status(module_name, import_module=import_module)

    monkeypatch.setattr(launcher, "_module_status", fake_status)
    assert "openpyxl" in launcher.missing_core_modules()
    assert "openpyxl" in launcher.missing_core_packages()
    exit_code = launcher.main(["--check-only", "--skip-port-check"])
    output = capsys.readouterr().out
    assert exit_code != 0
    assert "openpyxl" in output
    assert "repair_environment.bat" in output


def test_repair_script_only_installs_reported_missing_packages():
    text = (PROJECT_ROOT / "repair_environment.bat").read_text(
        encoding="utf-8"
    ).casefold()
    assert "--list-missing-core" in text
    assert "pip install --disable-pip-version-check %missing_packages%" in text
    assert "requirements-lock" not in text
    assert "pip install -r" not in text
    assert " -m venv" not in text


def test_qa_excel_can_be_read_and_replaced_atomically(tmp_path: Path):
    labels_path = tmp_path / "人工标签.xlsx"
    workbook = Workbook()
    workbook.active.title = "人工标注"
    workbook["人工标注"].append(["document_id", "正确标题", "reviewer_type"])
    workbook["人工标注"].append(["DOC-1", "原标题", "codex_agent"])
    workbook.save(labels_path)
    workbook.close()

    temporary = _write_updated_workbook(
        labels_path,
        "DOC-1",
        {"正确标题": "人工核对标题", "reviewer_type": "human_user"},
    )
    assert temporary.exists()
    original = load_workbook(labels_path, read_only=True, data_only=True)
    assert original["人工标注"]["B2"].value == "原标题"
    original.close()

    temporary.replace(labels_path)
    updated = load_workbook(labels_path, read_only=True, data_only=True)
    assert updated["人工标注"]["B2"].value == "人工核对标题"
    assert updated["人工标注"]["C2"].value == "human_user"
    updated.close()


def test_qa_context_reads_excel_labels(tmp_path: Path):
    db_path = tmp_path / "qa.db"
    data_root = tmp_path / "data"
    workspace = create_workspace({"workspace_name": "QA依赖回归"}, db_path, data_root)
    initialize_database(db_path)
    labels_path = tmp_path / "labels.xlsx"
    workbook = Workbook()
    workbook.active.title = "人工标注"
    workbook["人工标注"].append(["document_id", "正确标题", "reviewer_type"])
    workbook["人工标注"].append(["DOC-1", "测试标题", "codex_agent"])
    workbook.save(labels_path)
    workbook.close()

    context = qa_context(labels_path=labels_path, qa_db_path=db_path)
    assert context is not None
    assert context["workspace_id"] == workspace["workspace_id"]
    assert context["records"][0]["正确标题"] == "测试标题"


def test_home_gracefully_reports_missing_openpyxl_and_keeps_other_regions(
    tmp_path: Path,
    monkeypatch,
):
    from streamlit.testing.v1 import AppTest

    db_path = tmp_path / "home.db"
    data_root = tmp_path / "home_data"
    env_path = tmp_path / ".env"
    monkeypatch.setenv("PORTSCOPE_DB_PATH", str(db_path))
    monkeypatch.setenv("PORTSCOPE_DATA_ROOT", str(data_root))
    monkeypatch.setenv("PORTSCOPE_ENV_PATH", str(env_path))
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    monkeypatch.setattr(ui_daily, "_qa_excel_available", lambda: False)

    app = AppTest.from_file(str(PROJECT_ROOT / "app.py"), default_timeout=30).run()
    assert not app.exception
    warnings = [str(item.value) for item in app.warning]
    assert any("openpyxl" in message for message in warnings)
    assert any("今日工作台" in header.value for header in app.header)
    assert any(metric.label == "今日新增文档" for metric in app.metric)
    assert any(button.label == "开始采集" for button in app.button)
    assert not any("没有QA数据" in message for message in warnings)

    navigation = next(
        radio for radio in app.radio if "工作台" in list(radio.options)
    )
    navigation.set_value("人工审核").run()
    assert not app.exception
    assert any("openpyxl" in str(item.value) for item in app.warning)
    qa_scope = next(
        radio for radio in app.radio if "隔离QA人工金标准" in list(radio.options)
    )
    qa_scope.set_value("隔离QA人工金标准").run()
    assert not app.exception
    assert any("openpyxl" in str(item.value) for item in app.error)


def test_requirements_and_lock_contain_verified_openpyxl():
    requirements = (PROJECT_ROOT / "requirements.txt").read_text(encoding="utf-8")
    lock = (PROJECT_ROOT / "requirements-lock.txt").read_text(encoding="utf-8")
    assert "openpyxl>=3.1,<4.0" in requirements
    assert "openpyxl==3.1.5" in lock
    assert "et-xmlfile==2.0.0" in lock


def test_all_direct_runtime_dependencies_are_declared():
    rows = audit_runtime_dependencies(PROJECT_ROOT)
    assert rows
    missing = [row.package_name for row in rows if not row.declared]
    assert missing == []
    openpyxl_row = next(row for row in rows if row.import_module == "openpyxl")
    assert openpyxl_row.first_usage_file == "qa_review_store.py"
    assert openpyxl_row.core_dependency is True


def test_project_environment_imports_all_declared_direct_runtime_dependencies():
    if not PROJECT_PYTHON.exists():
        return
    modules = [
        row.import_module
        for row in audit_runtime_dependencies(PROJECT_ROOT)
    ]
    command = (
        "import importlib;"
        f"mods={modules!r};"
        "[importlib.import_module(name) for name in mods];"
        "print('DIRECT_RUNTIME_IMPORTS_OK')"
    )
    result = subprocess.run(
        [str(PROJECT_PYTHON), "-c", command],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "DIRECT_RUNTIME_IMPORTS_OK" in result.stdout


def test_daily_run_chain_still_has_zero_pip():
    content = "\n".join(
        (
            (PROJECT_ROOT / "run.bat").read_text(encoding="utf-8"),
            (PROJECT_ROOT / "scripts" / "launch_portscope.py").read_text(
                encoding="utf-8"
            ),
        )
    ).casefold()
    assert "pip install" not in content
    assert "requirements-lock" not in content
